"""Warm-list importer: Excel files on Drive (or a local file) -> persons.

Only files marked `consent_confirmed = true` in config are imported (docs/DESIGN.md section 16).
Nothing is ever written back to the Excel files.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

from .. import db, suppression
from ..identity import resolve_person

ALIASES = {
    "email": ["email", "e-mail", "email address", "work email", "personal email", "mail"],
    "name": ["name", "full name", "contact", "contact name"],
    "first": ["first name", "first", "given name", "firstname"],
    "last": ["last name", "last", "surname", "family name", "lastname"],
    "company": ["company", "organization", "organisation", "employer", "startup", "company name"],
    "title": ["title", "role", "job title", "position", "designation"],
    "location": ["location", "city", "state", "country", "region", "based in"],
    "linkedin": ["linkedin", "linkedin url", "linkedin profile", "profile", "linkedin link"],
}

US_STATES = {
    "alabama", "alaska", "arizona", "arkansas", "california", "colorado", "connecticut", "delaware",
    "florida", "georgia", "hawaii", "idaho", "illinois", "indiana", "iowa", "kansas", "kentucky",
    "louisiana", "maine", "maryland", "massachusetts", "michigan", "minnesota", "mississippi",
    "missouri", "montana", "nebraska", "nevada", "new hampshire", "new jersey", "new mexico",
    "new york", "north carolina", "north dakota", "ohio", "oklahoma", "oregon", "pennsylvania",
    "rhode island", "south carolina", "south dakota", "tennessee", "texas", "utah", "vermont",
    "virginia", "washington", "west virginia", "wisconsin", "wyoming", "district of columbia",
}
US_ABBR = {"al", "ak", "az", "ar", "ca", "co", "ct", "de", "fl", "ga", "hi", "id", "il", "in", "ia", "ks",
           "ky", "la", "me", "md", "ma", "mi", "mn", "ms", "mo", "mt", "ne", "nv", "nh", "nj", "nm", "ny",
           "nc", "nd", "oh", "ok", "or", "pa", "ri", "sc", "sd", "tn", "tx", "ut", "vt", "va", "wa", "wv",
           "wi", "wy", "dc"}
US_TOKENS = {"usa", "us", "u.s.", "u.s.a.", "united states", "united states of america", "america",
             "sf", "nyc", "bay area", "silicon valley", "san francisco", "seattle", "boston", "austin",
             "los angeles", "new york city", "chicago", "palo alto", "mountain view", "menlo park",
             "san jose", "sunnyvale", "cambridge, ma", "brooklyn", "berkeley", "redmond", "denver"}
NON_US_HINTS = {"canada", "uk", "united kingdom", "england", "india", "germany", "france", "london",
                "toronto", "vancouver", "berlin", "bangalore", "bengaluru", "singapore", "australia",
                "ireland", "netherlands", "spain", "italy", "israel", "china", "japan", "brazil", "mexico",
                "switzerland", "sweden", "europe", "eu"}


def us_located(location: str | None) -> int | None:
    """1 = US, 0 = clearly not US, None = unknown."""
    if not location:
        return None
    loc = location.strip().lower()
    parts = {p.strip(" .") for p in re.split(r"[,/|;()-]", loc) if p.strip()}
    if parts & NON_US_HINTS or any(h in loc for h in ("united kingdom", "canada", "india")):
        return 0
    if loc in US_TOKENS or parts & US_TOKENS or parts & US_STATES or parts & US_ABBR:
        return 1
    if any(t in loc for t in US_TOKENS if len(t) > 4) or any(s in loc for s in US_STATES):
        return 1
    return None


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " "))


def map_columns(header: list) -> dict[str, int]:
    cols: dict[str, int] = {}
    normed = [_norm_header(h) for h in header]
    for key, aliases in ALIASES.items():
        for i, h in enumerate(normed):
            if h in aliases and key not in cols:
                cols[key] = i
    if "email" not in cols:  # looser match, e.g. "Email (work)"
        for i, h in enumerate(normed):
            if "email" in h or "e-mail" in h:
                cols["email"] = i
                break
    return cols


@dataclass
class ImportResult:
    rows: int = 0
    imported: int = 0
    existing: int = 0
    invalid_email: int = 0
    suppressed: int = 0
    not_us: int = 0
    unknown_location: int = 0
    notes: list[str] = field(default_factory=list)


def read_rows(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out: list[dict] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_idx, cols = None, {}
        for i, row in enumerate(rows[:10]):
            c = map_columns(list(row))
            if "email" in c and len(c) >= 2:
                header_idx, cols = i, c
                break
        if header_idx is None:
            continue
        for row in rows[header_idx + 1:]:
            rec = {k: (str(row[i]).strip() if i < len(row) and row[i] is not None else "")
                   for k, i in cols.items()}
            if any(rec.values()):
                out.append(rec)
    return out


def _name(rec: dict) -> tuple[str, str]:
    full = rec.get("name") or " ".join(x for x in (rec.get("first"), rec.get("last")) if x)
    first = rec.get("first") or (full.split()[0] if full else "")
    return full.strip(), first.strip()


def import_rows(con, rows: list[dict], label: str, assume_us: bool = False,
                us_only: bool = True) -> ImportResult:
    res = ImportResult(rows=len(rows))
    for rec in rows:
        email = suppression.normalize_email(rec.get("email"))
        if not email:
            res.invalid_email += 1
            continue
        if suppression.is_suppressed(con, email):
            res.suppressed += 1
            continue
        loc = us_located(rec.get("location"))
        if loc is None and assume_us:
            loc = 1
        if us_only and loc == 0:
            res.not_us += 1
            continue
        if loc is None:
            res.unknown_location += 1
        full, _first = _name(rec)
        existing = db.find_person_by_identifier(con, "email", email)
        ids = [("email", email)]
        if rec.get("linkedin"):
            ids.append(("linkedin", rec["linkedin"]))
        pid = resolve_person(con, ids, full or email.split("@")[0], segment="warm", group_source="warm",
                             primary_employer=rec.get("company") or None, title=rec.get("title") or None,
                             country=rec.get("location") or None, us_located=loc, email=email,
                             email_source=f"warm:{label}", email_status="on_file")
        db.upsert_source_record(con, f"warm:{label}", email, "person", rec)
        if existing:
            res.existing += 1
        else:
            res.imported += 1
            _ = pid
    con.commit()
    return res


def import_drive_files(ctx, run) -> list[tuple[str, ImportResult]]:
    files = ctx.settings.get("drive.warm_files", []) or []
    us_only = bool(ctx.settings.get("pilot.us_only", True))
    results = []
    for f in files:
        label = f.get("label") or f["id"]
        if not f.get("consent_confirmed"):
            run.note(f"skipped warm file '{label}': consent_confirmed is not true")
            run.skipped += 1
            continue
        content = ctx.drive.download_xlsx(f["id"])
        res = import_rows(ctx.con, read_rows(content), label, bool(f.get("assume_us")), us_only)
        run.ok += res.imported
        results.append((label, res))
    return results
